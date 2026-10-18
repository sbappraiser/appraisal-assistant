"""Build src/a1.js from Fannie Mae Appendix A-1 (UAD 3.6 delivery spec): per 'PARENT/DataPoint' (or 'PARENT/@Attr')
the format type, max length and allowed enumerations.  usage: python3 -I make_a1.py <appendix-a-1.xlsx> <out a1.js>"""
import sys, json, warnings, collections
warnings.filterwarnings('ignore')
import openpyxl
wb = openpyxl.load_workbook(sys.argv[1], read_only=True)
spec = list(wb['UAD Delivery Spec 1.4'].iter_rows(values_only=True)); H = spec[0]; ix = {h: i for i, h in enumerate(H)}
en = list(wb['UAD Enumerations'].iter_rows(values_only=True)); EH = en[0]; ex = {h: i for i, h in enumerate(EH)}
enums = collections.defaultdict(lambda: collections.defaultdict(set))
for r in en[1:]:
    uid, name, val = r[ex['Unique ID']], r[ex['MISMO Data Point / Attribute Name']], r[ex['Supported Enumerations']]
    if uid and name and val not in (None, ''): enums[str(uid).strip()][str(name).strip()].add(str(val).strip())
out = {}
roles = ['Subject Property', 'Land Comparable', 'Property Analyzed Not Used', 'Rental Comparable', 'Gross Rent Multiplier Comparable', 'Sales Comparable']
for r in spec[1:]:
    parent = (r[ix['MISMO Parent Container']] or '').strip(); dp = (r[ix['MISMO Data Point Name']] or '').strip(); at = (r[ix['MISMO Attribute Name']] or '').strip()
    if not parent: continue
    names = ([dp] if dp else []) + [x.strip() for x in at.split('\n') if x.strip()]   # an attribute sits on the data point when there is one, else on the container
    if not names: out.setdefault(parent + '/', {'c': 1}); continue      # a container
    t = (r[ix['UAD Data Point Format Type']] or '').strip(); det = str(r[ix['UAD Data Point Format Details']] or '').strip()
    uid = str(r[ix['Unique ID']] or '').strip()
    fixedAll = str(r[ix['Supported Attribute Enumerations\n\n(Reference the MISMO Attribute Name column)']] or '').strip()
    for name in names:
        key = (dp + '/' + name if dp and name.startswith('@') else parent + '/' + name); o = out.setdefault(key, {})
        if not name.startswith('@'):
            if t: o['t'] = t
            if t == 'String' and det.isdigit(): o['len'] = max(o.get('len', 0), int(det))
            if t in ('Numeric', 'Amount', 'Date', 'Datetime') and det and det != 'None': o['f'] = det
        ev = set(enums.get(uid, {}).get(name, set()))
        if not ev and name.startswith('@'):   # "@ValuationUseType - SalesComparable" written in the row itself
            for line in fixedAll.split('\n'):
                if line.strip().startswith(name + ' - '): ev.add(line.split(' - ', 1)[1].strip())
            if not ev and len(names) == 1 and fixedAll and ' - ' not in fixedAll: ev = {fixedAll}
        if not ev and t == 'Enumerated' and not name.startswith('@') and uid in enums: ev = set().union(*enums[uid].values())
        if name.startswith('@xlink'): ev = set()
        if ev: o['e'] = sorted(set(o.get('e', [])) | ev)
        if r[ix['Conditionality']] == 'R': o['r'] = 1
open(sys.argv[2], 'w').write('/* Fannie Mae Appendix A-1 (UAD 3.6 delivery spec v1.4), made by tools/make_a1.py – allowed values and formats per PARENT/DataPoint */\nconst A1=' + json.dumps(out, separators=(',', ':')) + ';\n')
print(len(out), 'entries', sum(1 for v in out.values() if 'e' in v), 'enumerated')
